# Builds the prod import workbook (Kissflow-internal) from the dev export: one sheet per form, headed by field names.
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
D = "/Users/thomas/Documents/Kissflow Projects/POC/Beiersdorf/kissflow-build/night-2026-10-21/prod-seed/"
OUT = "/Users/thomas/Documents/Kissflow Projects/POC/Beiersdorf/kissflow-build/CIQ-prod-seed-INTERNAL.xlsx"
CFG = [("Config_Key", "Config Key"), ("Config_Kind", "Config Kind"), ("Version_Label", "Version Label"), ("Config_State", "Config State"), ("Section_Key", "Section Key"),
       ("Section_Rev", "Section Rev"), ("Payload_JSON", "Payload JSON"), ("Internal_Annotations_JSON", "Internal Annotations JSON"), ("Content_Hash", "Content Hash"),
       ("Evaluator_Version", "Evaluator Version"), ("Change_Note", "Change Note"), ("Published_At", "Published At"), ("Release_Key", "Release Key"), ("Release_Label", "Release Label"),
       ("Release_Stage", "Release Stage"), ("Release_Manifest_JSON", "Release Manifest JSON"), ("Release_Bundle_JSON", "Release Bundle JSON"), ("Manifest_Hash", "Manifest Hash"),
       ("Regression_Result_JSON", "Regression Result JSON"), ("Test_Case_Codes", "Test Case Codes")]
HELP = [("Article_Title", "Article Title"), ("Article_Body", "Article Body"), ("Sort_Order", "Sort Order")]
F = "Arial"; H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F3864")
wb = Workbook(); rm = wb.active; rm.title = "Read me"
notes = [
  "CIQ prod seed: Kissflow-internal. Contains the webhook tokens (SET:current); do not share.",
  "",
  "Data is not pushed with the app. Import these into PRODUCTION after the push, as the business administrator:",
  "1. App CIQ Digital > form 'CIQ Config Version' > Import: the sheet 'CIQ Config Version' (12 rows; headers are the field names).",
  "2. Supplier portal > form 'CIQP Help Article' > Import: the sheet 'CIQP Help Article' (9 FAQ rows; the dev row 'Test' is left out).",
  "If the import reads only the first sheet, use the one-sheet files with the same rows: CIQ-prod-seed-config-version-INTERNAL.xlsx and CIQ-prod-seed-help-article-INTERNAL.xlsx.",
  "3. Every text cell must stay exactly as it is (the JSON and the hashes are checked by the app). The biggest cell is about 25,600 characters.",
  "4. 'Published By' (a user field) is left out: dev users are not prod users. The app does not need it to run.",
  "5. SET:current carries the dev webhook URLs. The screens build the prod URLs at runtime (prod domain + prod account + the same token), so the row is imported unchanged.",
  "6. Test suppliers are created in prod by the business administrator in the app (Suppliers > Add supplier), as test case TC-01 does.",
  "If an import refuses a long cell, tell me: I give you a small script that creates the rows over REST with a prod key you enter yourself.",
]
for i, t in enumerate(notes, 1):
  c = rm.cell(row=i, column=1, value=t); c.font = Font(name=F, bold=(i == 1), size=11 if i == 1 else 10); c.alignment = Alignment(wrap_text=True)
rm.column_dimensions["A"].width = 130
def sheet(name, cols, rows):
  ws = wb.create_sheet(name)
  for j, (_, label) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=label); c.font = H; c.fill = HF; ws.column_dimensions[c.column_letter].width = 22
  for i, r in enumerate(rows, 2):
    for j, (k, _) in enumerate(cols, 1):
      v = r.get(k)
      if isinstance(v, (dict, list)): v = json.dumps(v)
      cell = ws.cell(row=i, column=j, value=v)
      cell.font = Font(name=F, size=9)
      if isinstance(v, str): cell.number_format = "@"
  ws.freeze_panes = "A2"
  return len(rows)
n1 = sheet("CIQ Config Version", CFG, json.load(open(D + "config-version.json")))
# 22 Oct review: two dev rows differ from the portal's live FAQ text (prototype/src/features/supplier/SupplierHelp.tsx
# FAQS_LIVE): row 1's title has a stray "??", and row 9 still says answers can change "until you give the final release",
# while the portal locks the questionnaire once the release is requested. Both take the live text. Rows 7-8 (e-mails)
# wait for Thomas's decision D1 (MORNING-2026-10-22.md).
FAQ_FIX = {
  "Why do I only see some sections??": ("Article_Title", "Why do I only see some sections?"),
  "Can I change an answer after confirming completeness?": ("Article_Body", "Yes, until Beiersdorf requests your final release: correct data and add documents directly in the questionnaire. Beiersdorf's reviewers are told about each change and look at it again, and a changed metadata answer can then only add sections. Once your final release is requested the questionnaire is locked: if something must change, contact Beiersdorf, who can reopen the review. After the release the questionnaire is frozen; to report a change, use \u201cReport a change\u201d on the released request."),
}
def help_rows():
  rows = json.load(open(D + "help-article.json"))
  for r in rows:
    fix = FAQ_FIX.get(r.get("Article_Title"))
    if fix: r[fix[0]] = fix[1]
  return rows
n2 = sheet("CIQP Help Article", HELP, help_rows())
wb.save(OUT); print("saved", OUT, n1, "config rows,", n2, "help articles")
# 22 Oct review: an importer may read only the first sheet (the Read me), so each form also gets a one-sheet file
KB = "/Users/thomas/Documents/Kissflow Projects/POC/Beiersdorf/kissflow-build/"
for name, cols, src, out in [("CIQ Config Version", CFG, "config-version.json", "CIQ-prod-seed-config-version-INTERNAL.xlsx"),
                             ("CIQP Help Article", HELP, "help-article.json", "CIQ-prod-seed-help-article-INTERNAL.xlsx")]:
  wb = Workbook(); wb.remove(wb.active)
  n = sheet(name, cols, help_rows() if src == "help-article.json" else json.load(open(D + src)))
  wb.save(KB + out); print("saved", KB + out, n, "rows")

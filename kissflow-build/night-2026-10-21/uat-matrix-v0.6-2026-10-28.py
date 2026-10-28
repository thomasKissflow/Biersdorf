# Updates the client workbook CIQ-UAT-test-cases.xlsx to the matrix v0.6 build of 28 Oct (MATRIX-v0.6-ASSUMPTIONS-
# 2026-10-28.md): SF-02 and SF-08 dropped (9 numbered metadata questions, 8 required), the K2.1 component rules (labels,
# Additive Yes/No, CAS/EC/REACH checks, the D6 origin check), the Feedstock fields in K2.3, the read-only Annex 1
# country of origin, no "Cannot provide" on the three controller questions, Sustainability reviewing only through
# EUDR (K6.17), and Test rules without the palm question. Apply it in the same push as the UIs and the configuration
# rows of 28 Oct. Dry run by default (shows every change and checks each anchor occurs exactly once); --apply writes
# the workbook after copying it to backups/.
import shutil, sys
import openpyxl

KB = "/Users/thomas/Documents/Kissflow Projects/POC/Beiersdorf/kissflow-build/"
APPLY = "--apply" in sys.argv

TC09_STEPS = """1. Open K2.1 Components and additives and click “Add component”.
2. Fill “INCI (European) / IUPAC” and “Function”, answer “Additive?” with No, and open the “Origin” list. Keep “Vegetable (bio mass)”.
3. Enter CAS number 59-02-8, then 59-02-9. Enter EC number 200-412-3, then 200-412-2, then replace it with the REACH registration number 01-2119457892-27-0000.
4. Enter the concentration as 60 – 40 % and click outside the field, then enter 40 – 60 %.
5. Add a second component: choose Origin “Petrochemical / fossil”, fill INCI and Function, answer “Additive?” with Yes and enter 0.1 – 0.5 %.
6. Add a third component. Click into its “INCI (European) / IUPAC” field, leave everything empty and click elsewhere. Click “Check completeness”. Then remove the third component with its bin icon.
7. In Metadata questions, deselect “Petrochemical / fossil”. Go back to K2.1 and click “Check completeness”. Then select “Petrochemical / fossil” again.
8. Enter “Natural origin index” 1.5, then 0.853, then 0.85.
9. Reload the page.
10. As RMM, open K2.1 on the case's Questionnaire tab."""

TC09_EXPECTED = """- “Composition*” reads “One row per component. Each component's origin must be one you selected in the metadata questions, and each selected origin must be used by a component.” Before the first row: “No components yet. Add every component of the raw material, including additives such as antioxidants.”
- Each “Component n” has “Additive?*” with Yes and No (neither preselected), “INCI (European) / IUPAC*” with the note “The INCI name. If there is none: the chemical name (IUPAC). For colorants, also the Colour Index number.”, “CAS number”, “EC number (EINECS/ELINCS/NLP) or REACH registration no.”, “Function*”, “Origin*” and “Concentration in the raw material*” (min – max %).
- The Origin list offers only Vegetable (bio mass) and Petrochemical / fossil.
- Step 3: at once, 59-02-8 gives “The CAS check digit does not match. Please check for a typo” and 200-412-3 gives “The EC check digit does not match. Please check for a typo”. The valid numbers are accepted, and so is the REACH registration number in the EC field.
- Step 4: 60 – 40 gives “The minimum is higher than the maximum”. Values outside 0–100 give “Concentrations must be between 0 and 100 %”.
- “Total: min x % · max y %” is shown. When the maximums add up to less than 100 %: “The maximum concentrations add up to y %, less than 100 %. Is a component missing? This is a hint only and does not block your submission.”
- Step 6: the empty component shows “Required” under INCI (European) / IUPAC, Function and Additive?, and “Enter a minimum and a maximum concentration”. The completeness check lists “Component 3: INCI (European) / IUPAC is missing”, “Component 3: Function is missing”, “Component 3: Additive: answer Yes or No” and “Component 3: Enter a minimum and a maximum concentration”; each item jumps to the component. A wrong CAS or EC number is listed the same way. After the removal the row disappears; the “<n> components” badge and “Last changed by <name> · <time>” update.
- Step 7: component 2 shows “Not one of the origins you selected in the metadata questions” under Origin, and the completeness check lists “Component 2: its origin “Petrochemical / fossil” is not one of the origins you selected in the metadata questions”. After selecting the origin again, both are gone. (TC-07 checks the other direction: a selected origin that no component has.)
- Step 8: 1.5 gives “Must be between 0 and 1”; 0.853 gives “Max. 2 decimals”; 0.85 is accepted.
- Step 9: all rows and values come back in the same order.
- Step 10: RMM sees the same rows read-only (#, INCI / IUPAC, CAS, EC / REACH no., Function, Origin, Additive, Concentration). An unanswered Additive shows “—”.
- Later, a component Beiersdorf has asked about shows “Asked about” instead of the bin and cannot be removed."""

WORKBOOK = {  # sheet "Test cases": cell → [(old, new)]; old must occur exactly once in the cell
    # TC-05: 8 required metadata questions (SF-08 dropped)
    "G6": [("(0 of 9 answered)", "(0 of 8 answered)"), ("metadata questions: 0 of 9 answered", "metadata questions: 0 of 8 answered")],
    # TC-07: 9 numbered questions (SF-02, SF-08 dropped), 8 required; Feedstock in K2.3; A1's country of origin; D6
    "G8": [
        ("- 11 numbered questions, 9 of them required (marked *). “Target markets set by Beiersdorf” is read-only. The header reads “0 of 9 answered”.",
         "- 9 numbered questions, 8 of them required (marked *). “Target markets set by Beiersdorf” is read-only. The header reads “0 of 8 answered”."),
        ("and K2.3 shows “Part(s) used” and “Country of origin”. The panel lists the added annexes with “Because: …” and the evidence that will be asked for.",
         "and K2.3 shows “Part(s) used” and “Country of origin”. (K2.3 asks for “Feedstock (e.g. sunflower, palm, soy)”, “Content [%] in INCI” and “Content [%] in raw material” on every case, whatever the origin.) The panel lists the added annexes with “Because: …” and the evidence that will be asked for.\n- Step 3: A1 starts with “Country of origin”, locked: “Filled in from Country of origin (K2.3). Change it there.”"),
        ("You get the standard sections; Beiersdorf may ask for more.” The supplier does not have to wait for Beiersdorf.",
         "You get the standard sections; Beiersdorf may ask for more.” The supplier does not have to wait for Beiersdorf. Each selected origin must be used by at least one component in K2.1: until a component has “Waste / side stream”, the completeness check lists “You selected the origin “Waste / side stream”, but no component has it. Add the component or remove the origin.”"),
        ("“Required metadata questions answered: 9 of 9”", "“Required metadata questions answered: 8 of 8”"),
    ],
    # TC-08: no "Cannot provide" on the three controller questions (D10)
    "G9": [("- Step 5: Yes shows “Pesticides analysis (PDF)”; No shows “Reason the pesticides analysis is not available” instead.",
            "- Step 5: Yes shows “Pesticides analysis (PDF)”; No shows “Reason the pesticides analysis is not available” instead. This question offers no “Cannot provide? Add a justification”, and neither do “Is palm (kernel) oil used as feedstock?” (A1) and “Kinematic viscosity at 100 °C: availability” (A4).")],
    # TC-12: 8 required metadata questions
    "G13": [("“Required metadata questions answered: x of 9”", "“Required metadata questions answered: x of 8”")],
    # TC-15: the reviewer's K2.1 table
    "G16": [("- K2.1 shows one row per component (INCI, CAS, EC, origin, concentration), each with “Ask supplier”.",
             "- K2.1 shows one row per component (INCI / IUPAC with function and additive, CAS / EC or REACH no., origin, concentration), each with “Ask supplier”.")],
    # TC-17: A1's derived country of origin; the D6 origin check after a Part 2 deselect (it would block the release)
    "F18": [("4. Deselect “Vegetable (bio mass)” in the origin question.",
             "4. Deselect “Vegetable (bio mass)” in the origin question and read the message. Look at K2.1, then select “Vegetable (bio mass)” again.")],
    "G18": [
        ("- Step 2: the change is saved with your name and time.",
         "- Step 2: the change is saved with your name and time; A1's locked “Country of origin” shows the new value."),
        ("You confirmed completeness, so sections are only added now.” No section disappears.",
         "You confirmed completeness, so sections are only added now.” No section disappears. While Vegetable is deselected, each component with that origin shows “Not one of the origins you selected in the metadata questions” under Origin (it would block the final release); selecting Vegetable again clears it."),
    ],
    # TC-30: Sustainability reviews only through EUDR (K6.17); Annex 1 is led by Regulatory Affairs
    "C31": [("Quality Management reviewer (repeat for Microbiology, Sustainability and Regulatory Affairs)",
             "Quality Management reviewer (repeat for Microbiology and Regulatory Affairs, and for Sustainability on a case with the EUDR question answered Yes)")],
    # TC-44: the palm question SF-08 and its RSPO evidence (R-80) are gone
    "F45": [("Set the origin to Vegetable and “Is any feedstock palm-based (palm or palm kernel oil)?” to Yes.", "Set the origin to Vegetable.")],
    "G45": [("- Step 2: “Derived scope” lists A1 Vegetable origin as mandatory; “Evidence the supplier must upload” includes the RSPO evidence;",
             "- Step 2: “Derived scope” lists A1 Vegetable origin, A2 Manufacturing process and A9 Polyaromatic hydrocarbons (PAH) as mandatory; “Evidence the supplier must upload” lists “Pesticide and mycotoxin analysis” and “Sustainability certificates” (R-10) besides the documents every case needs;")],
}
# TC-09: rewritten as a whole (labels, Additive Yes/No, REACH number, completeness items, the D6 origin check)
WHOLE = {
    "F10": ("tick “Additive” and enter 0.1 – 0.5 %.", TC09_STEPS),
    "G10": ("Each “Component n” has INCI name*, CAS number, EC number, Function*, Origin*, Concentration in the raw material* (min – max %) and Additive.", TC09_EXPECTED),
}

ok = True
wb = openpyxl.load_workbook(KB + "CIQ-UAT-test-cases.xlsx")
ws = wb["Test cases"]
for cell, pairs in WORKBOOK.items():
    v = ws[cell].value or ""
    for old, new in pairs:
        n = v.count(old)
        print(f"workbook {cell} ({ws.cell(ws[cell].row, 1).value}): {'OK ' if n == 1 else 'MISS'} {old[:70]}…")
        ok &= n == 1
        v = v.replace(old, new)
    ws[cell].value = v
for cell, (anchor, new) in WHOLE.items():
    v = ws[cell].value or ""
    n = v.count(anchor)
    print(f"workbook {cell} ({ws.cell(ws[cell].row, 1).value}): {'OK ' if n == 1 else 'MISS'} whole cell (anchor {anchor[:50]}…)")
    ok &= n == 1
    ws[cell].value = new
if not ok:
    sys.exit("an anchor did not match exactly once: nothing written")
if APPLY:
    shutil.copy(KB + "CIQ-UAT-test-cases.xlsx", KB + "night-2026-10-21/backups/CIQ-UAT-test-cases.before-matrix-v0.6.xlsx")
    wb.save(KB + "CIQ-UAT-test-cases.xlsx")
    print("applied (backup in night-2026-10-21/backups/)")
else:
    print("dry run: every anchor matches; run with --apply")

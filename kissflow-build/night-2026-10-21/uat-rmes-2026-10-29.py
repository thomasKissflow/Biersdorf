# Updates the client workbook CIQ-UAT-test-cases.xlsx to the RMES of BDF's e-mail of 28 Oct ("exactly the scope … no
# more"): the RMES is one read-only sheet of 8 fields (Raw material / tradename, Supplier, Manufacturer; per component
# INCI, %, CAS #, EC #, Source of origin and its free text), filled automatically from the CIQ, each with its CIQ source,
# shown to Safety and Regulatory only. The criteria (Parts A–D), limit checks, the suggestion box, findings and the RMM
# "RMES progress" card are gone. The Safety / Regulatory Approved/Rejected decision stays as the review's decision,
# outside the RMES, on the Questionnaire tab's decision card.
# - TC-26 to TC-29 are rewritten (TC-29, the Rejected flow, keeps its course with the new navigation).
# - Every other RMES / criteria / finding mention is fixed: TC-02, 14, 18, 30, 31, 33, 40; and the REQ-RMES-* ids of
#   deferred RMES requirements leave the Requirement column of TC-16, 17, 24, 41 and 42. The K2.2 "BDF limit" of TC-08
#   (the supplier's heavy-metal limits) is not RMES and stays; "Comments / findings" is the testers' column.
# Client-facing: no internal file is named. Formatting, dropdowns and the Summary formulas stay (cell values only; the
# row heights of rewritten rows grow to fit). Dry run by default (shows every change and checks each anchor occurs
# exactly once, and that no RMES criteria / finding wording is left); --apply writes the workbook. The backup
# backups/CIQ-UAT-test-cases.before-rmes-2026-10-29.xlsx is made once, before the first write, and never overwritten.
import math, os, re, shutil, sys
import openpyxl

KB = "/Users/thomas/Documents/Kissflow Projects/POC/Beiersdorf/kissflow-build/"
BOOK = KB + "CIQ-UAT-test-cases.xlsx"
BACKUP = KB + "night-2026-10-21/backups/CIQ-UAT-test-cases.before-rmes-2026-10-29.xlsx"
APPLY = "--apply" in sys.argv

# ---------------------------------------------------------------------------------------------- TC-26 to TC-29
TC26 = {
    "B": "RMES",
    "C": "Product Safety reviewer (repeat as Regulatory Affairs); RMM, Quality Management reviewer and Supplier contact (to check)",
    "D": "RMES shows the 8 fields filled automatically, with their CIQ sources (Safety and Regulatory)",
    "E": "A case in review with Safety and Regulatory tasks. In K2.1 the supplier entered at least three components: two of the same origin (e.g. Vegetable (bio mass)) and one of another origin, one of them without CAS and EC number. K2.3 holds the specification of each origin used (for Vegetable or Animal: Part(s) used and Country of origin). In General the supplier answered “Manufacturer (name and address)”.",
    "F": """1. As the Safety reviewer, open the case from My reviews and click the tab “RMES” next to “Questionnaire”.
2. Read the header: Raw material / tradename, Supplier and Manufacturer.
3. Click the source line under Manufacturer.
4. Go back to the RMES tab and read the Composition table and the line below it.
5. In the Free text column, click the source line of one text.
6. Look for anything on the RMES tab that you can type, choose or record.
7. As the Regulatory reviewer, open the same case's RMES tab.
8. As RMM and as the Quality Management reviewer, open the case's review workspace. As RMM, also open the case page > Overview.
9. As the supplier, open the request and its questions.""",
    "G": """- Step 1: the tab is named “RMES” (no count). It shows “Raw Material Evaluation Sheet” with “RMES · <case ID> · internal: never visible to the supplier · filled automatically from the CIQ” and “Record created <date, time>”.
- Step 2: exactly three header fields. Raw material / tradename shows the trade name, Supplier the supplier's name, and Manufacturer the supplier's answer to “Manufacturer (name and address)”, never the supplier's name in its place. Each shows its CIQ source, e.g. “General · FLD-GEN-004 Manufacturer (name and address)”, and “as of <date, time>”. An unanswered field reads “Not answered yet”; an answer the supplier could not provide reads “Cannot provide: <justification>”.
- Step 3: the Questionnaire tab opens at General.
- Step 4: one row per K2.1 component, in the same order, with the columns #, INCI, %, CAS #, EC #, Source of origin and Free text. Each column names its CIQ field (FLD-2.1-001, FLD-2.1-007, FLD-2.1-002, FLD-2.1-003, FLD-2.3-01). % shows the CIQ range, e.g. “90–99 %”, or one value when minimum and maximum are equal. A blank CAS # or EC # shows “—”; EC # may hold a REACH registration number. Source of origin shows the component's origin. Free text shows that origin's specification from K2.3 with its field ID, e.g. “FLD-2.3-11 Fossil specification (e.g. oil, wax, gas)”; for Vegetable and Animal “FLD-2.1-101 Part(s) used” and “FLD-2.3-10 Country of origin”. The rows of the same origin, and a Vegetable row with an Animal row (the CIQ asks Part(s) used and Country of origin once for both), show the same text, marked “*”, and the line below says “The CIQ asks the free text once per origin type, not per component, and asks vegetable and animal together: components of the same origin, or of vegetable and animal origin, show the same text.”
- Step 5: the Questionnaire tab opens at K2.3.
- Step 6: nothing on the RMES tab can be edited. There are no criteria, limit checks, suggestions, findings or decision, and no Process ID / RMT ID / IRMA ID or Declaration.
- Step 7: Regulatory sees the same sheet with the same values.
- Step 8: RMM and Quality Management have no RMES tab, and the case page shows no RMES card.
- Step 9: nothing of the RMES is visible to the supplier: not the sheet, its labels or the word “RMES”.
- If the tab shows “Kissflow did not let your role read this case's questionnaire, so the RMES can't be filled.”, the reviewer's role cannot read the case: record Blocked and name the role.""",
    "H": "REQ-PROC-021, REQ-PROC-022, REQ-CFG-004, REQ-RMES-015",
    "I": "High",
}
TC27 = {
    "B": "RMES",
    "C": "Supplier contact; Product Safety reviewer (to check)",
    "D": "RMES follows a supplier change in Part 2; unchanged after the final release",
    "E": "TC-26 done. The case is in review and the final release has not been requested yet.",
    "F": """1. As the Safety reviewer, open the case's RMES tab and note the values.
2. As the supplier, open the request. In General, change “Manufacturer (name and address)”. In K2.1, change one component's concentration and add a component of an origin already used. In K2.3, change the specification of one origin used.
3. As the Safety reviewer, click Reload and open the RMES tab again.
4. Complete the case: every review completes and the supplier gives the final release (TC-31, TC-32).
5. As the Safety reviewer, open the completed case from My reviews > Completed and its RMES tab.""",
    "G": """- Step 3: the sheet shows the new values without anyone copying them: the new manufacturer with a new “as of” time, the new range, the new component in its K2.1 position (marked “*” with the other rows of its origin), and the new specification in every row of that origin. The record line and “Record created” are unchanged. Nothing has to be confirmed or saved on the RMES.
- Step 5: the sheet shows exactly the released values. They cannot change any more, because the released CIQ is frozen.""",
    "H": "REQ-PROC-022, REQ-PROC-016, REQ-PROC-021",
    "I": "High",
}
TC28 = {
    "B": "Safety and Regulatory decision",
    "C": "Product Safety reviewer (repeat as Regulatory Affairs); Regulatory Affairs reviewer and RMM (to check); Supplier contact (to check)",
    "D": "Safety/Regulatory decision and completing the review",
    "E": "A case in review with a Safety task and its RMES record (TC-26). No Safety open point is open; no Safety decision is recorded yet.",
    "F": """1. In the review workspace, read the bottom bar and expand the list of blocking items.
2. Click “Record the decision” in the bar.
3. Read the “Safety decision” card. Choose “Approved”, add an optional comment and click “Record approval”.
4. Read the bottom bar. Click “Change decision”, then “Cancel”.
5. Open the RMES tab.
6. Click “Complete my review”.
7. As the Regulatory reviewer, open the same case's review workspace.
8. As RMM, open the case > Overview.
9. As the supplier, open the request.""",
    "G": """- Step 1: “Safety decision: not recorded yet” with “Record the decision” and “Internal: the supplier never sees it.” The bar reads “1 item blocks your completion”, and the list names “The Safety decision (Approved or Rejected) is not recorded yet”; clicking it goes to the decision card.
- Step 2: the Questionnaire tab opens with the card “Safety decision” at the top of the right column, above the open points: “Internal. Only Product Safety records it; nobody decides for another department. It is not part of the RMES.”
- Step 3: the card shows “Approved” with the comment, name and time. The bar shows “Safety decision: Approved” and “Internal: the supplier never sees it.”
- Step 4: the bar shows “Nothing blocks your review” and “… It is not a raw-material approval.”
- Step 5: the RMES tab shows only the 8 fields, not the decision.
- Step 6: the bar reads “Product Safety review completed” with name, e-mail, time and “not a raw-material approval”. The case moves to My reviews > Completed.
- Step 7: Regulatory has only its own “Regulatory decision” card; it cannot record or change the Safety decision.
- Step 8: Reviews shows Product Safety as “Completed” with “Decision: Approved”. Confirmations shows the review completion of Safety with name, e-mail and time.
- Step 9: the supplier sees no decision and gets no message for this completion.
- Before the case's RMES record exists, the card shows “Not possible yet” and “Record approval” is disabled.""",
    "H": "REQ-RMES-013, REQ-PROC-015, REQ-PROC-038, REQ-PROC-039",
    "I": "High",
}
TC29 = {  # the Rejected flow: unchanged course, the decision card now on the Questionnaire tab
    "B": "Safety and Regulatory decision",
    "E": "A case in review. The Safety reviewer has no open point and no decision on it yet. RMM has not completed its own review.",
    "F_old": "1. As the Safety reviewer, choose “Rejected” in the Safety decision card and click “Record rejection” without a reason.",
    "F_new": "1. As the Safety reviewer, open the case's review workspace. On the Questionnaire tab, choose “Rejected” in the “Safety decision” card and click “Record rejection” without a reason.",
}
WHOLE = {27: TC26, 28: TC27, 29: TC28}  # sheet rows
OLD_TITLE = {27: "Record RMES criterion results with source, limits and required comments", 28: "Add an internal finding and ask the supplier about it", 29: "Record the Safety decision and complete the Safety review", 30: "A Rejected Safety decision is flagged to RMM, who decides how the case goes on"}

# ---------------------------------------------------------------------------------------------- other mentions
EDITS = {  # cell → [(old, new)]; old must occur exactly once in the cell
    # TC-02: the reviewer page subtitles
    "G3": [("- Only the Safety and Regulatory page subtitles mention the internal RMES evaluation and its decision.",
            "- Only the Safety and Regulatory page subtitles mention their internal decision and the RMES tab (“The RMES tab shows the material's evaluation sheet, filled automatically from the CIQ.”).")],
    # TC-14: the RMES record of a case in review
    "G15": [("- Step 7: the RMES tab shows “Record created <date> for <case ID>”.",
             "- Step 7: the RMES tab shows “RMES · <case ID> · internal: never visible to the supplier · filled automatically from the CIQ” and “Record created <date, time>”, with the case's 8 fields (TC-26).")],
    # TC-16, TC-17, TC-24: REQ-RMES-011 (finding responses) and -012 (criteria reopening) are deferred with the criteria
    "H17": [("REQ-RMES-011, REQ-CTL-003", "REQ-CTL-003")],
    "H18": [(", REQ-RMES-012", "")],
    "H25": [("REQ-RMES-011, ", "")],
    # TC-18: the ask window's note and what the supplier sees
    "G19": [("Internal notes, RMES results and findings stay internal.", "Internal notes, the RMES and the Safety and Regulatory decisions stay internal."),
            ("No internal note, RMES result or finding is visible.", "No internal note, RMES content or decision is visible.")],
    # TC-30: Regulatory Affairs also needs its decision
    "G31": [("- Step 2: each item names its point ID,",
             "- For Regulatory Affairs the list also names its decision until it is recorded (TC-28).\n- Step 2: each item names its point ID,")],
    # TC-31, TC-33: what the supplier and the final document never show
    "G32": [("- No decision, finding, RMES result or internal note is shown to the supplier.", "- No decision, RMES content or internal note is shown to the supplier.")],
    "G34": [("It shows no RMES decision, finding or internal note, and cannot be changed.", "It shows no RMES content, no Safety or Regulatory decision and no internal note, and cannot be changed.")],
    # TC-40: no findings any more; the question is an ordinary one
    "E41": [("On R1, reviewers wrote internal notes and RMES results, raised a finding and turned it into a question.",
             "On R1, reviewers wrote internal notes, recorded the Safety and Regulatory decisions with a comment and asked the supplier a question.")],
    "F41": [("4. As A, open R1 and the question that came from the finding. Read the Help page.", "4. As A, open R1 and the question from Beiersdorf. Read the Help page.")],
    "G41": [("no RMES criteria or results, no findings, no internal notes,", "nothing of the RMES, no internal notes,"),
            ("The question from the finding appears as “Question from Beiersdorf”, with no sign of the finding.", "The question appears as “Question from Beiersdorf”, with its reference, question, justification and due date only.")],
    # TC-41, TC-42: REQ-RMES-019 (RMES notifications) and -018 (RMES progress per department) are deferred
    "H42": [(", REQ-RMES-019", "")],
    "H43": [(", REQ-RMES-018", "")],
}
# no RMES criteria / finding / limit wording may be left anywhere on the sheet (the testers' "Comments / findings"
# column head and the K2.2 supplier limit of TC-08 aside)
LEFT = re.compile(r"criteri|finding(?!s\b(?:,| column))|Parts? [A-D]\b|RMES result|RMES progress|RMES · Safety|limit box|AI suggestion|REQ-RMES-0(?:0[1-9]|1[012]|1[6-9]|20)", re.I)
ALLOWED = {"K1", "G27"}  # the "Comments / findings" column head; TC-26's step 6, which says the criteria are gone


def height_for(ws, r):
    """The row height (points) the wrapped text of Steps and Expected result needs, as the workbook's own rows are set
    (Arial 9, columns of 60: about 70 characters and 13 points a line)."""
    lines = max(sum(max(1, math.ceil(len(p) / 70)) for p in str(ws.cell(r, c).value or "").split("\n")) for c in (5, 6, 7))
    return max(ws.row_dimensions[r].height or 0, math.ceil(lines * 13 / 12.5) * 12.5)


ok = True
wb = openpyxl.load_workbook(BOOK)
ws = wb["Test cases"]
col = {ws.cell(1, c).value: openpyxl.utils.get_column_letter(c) for c in range(1, ws.max_column + 1)}
for r, title in OLD_TITLE.items():
    n = ws[f"D{r}"].value == title
    print(f"row {r} ({ws[f'A{r}'].value}): {'OK ' if n else 'MISS'} title “{title[:60]}”")
    ok &= n
for r, tc in WHOLE.items():
    for letter, v in tc.items():
        ws[f"{letter}{r}"].value = v
    print(f"row {r} ({ws[f'A{r}'].value}): rewritten as “{tc['D']}”")
r = 30
v = ws[f"F{r}"].value or ""
n = v.count(TC29["F_old"])
print(f"row {r} ({ws[f'A{r}'].value}): {'OK ' if n == 1 else 'MISS'} step 1 with the new navigation")
ok &= n == 1
ws[f"F{r}"].value = v.replace(TC29["F_old"], TC29["F_new"])
ws[f"B{r}"].value = TC29["B"]
ws[f"E{r}"].value = TC29["E"]
for cell, pairs in EDITS.items():
    v = ws[cell].value or ""
    for old, new in pairs:
        n = v.count(old)
        print(f"{cell} ({ws.cell(ws[cell].row, 1).value}): {'OK ' if n == 1 else 'MISS'} {old[:70]}…")
        ok &= n == 1
        v = v.replace(old, new)
    ws[cell].value = v
for r in [27, 28, 29, 30, 31]:
    ws.row_dimensions[r].height = height_for(ws, r)
left = [(c.coordinate, m.group(0)) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.coordinate not in ALLOWED for m in [LEFT.search(c.value)] if m]
print("left over RMES criteria / finding wording:", left or "none")
ok &= not left
if not ok:
    sys.exit("an anchor did not match exactly once, or RMES criteria wording is left: nothing written")
if APPLY:
    if not os.path.exists(BACKUP):
        shutil.copy(BOOK, BACKUP)
    wb.save(BOOK)
    print("applied (backup in night-2026-10-21/backups/CIQ-UAT-test-cases.before-rmes-2026-10-29.xlsx)")
else:
    print("dry run: every anchor matches; run with --apply")
